#!/usr/bin/env python3
"""Form 207 document register: read, update and check.

Form 207 (Document Transmittal) is kept on each project as the master record
of drawings and reports RECEIVED. One row per document, grouped in discipline
bands; one column per issue date; CURRENT ISSUE and STATUS at the right-hand end.

The script never edits the sheet cell by cell. It reads the register into a
plain data model, applies the changes, and re-renders every row below the date
header from that model, using the sheet's own cell styles. That is what keeps
bands, borders, merges and the print area consistent however many rows are
added -- and it is why a new drawing can be slotted into its band without
anyone inserting rows by hand.

    python3 register_207.py read     <register.xlsx>                  -> JSON model on stdout
    python3 register_207.py new      <form207.xlsx> <changes.json> <out.xlsx>   (first receipt)
    python3 register_207.py apply    <register.xlsx> <changes.json> <out.xlsx>
    python3 register_207.py validate <register.xlsx>

`apply` never writes over its input. The caller saves to a new filename and
moves the old file to _SS (see SKILL.md section 8).

changes.json:
{
  "header": {"date": "14 October 2026"},              # optional, any header key
  "placeholders": {"{{Recipient 1}}": "Not applicable"},  # optional, replaces form placeholders
  "received": [
    {"date": "2026-10-13",                            # title-block issue date
     "discipline": "ARCHITECTURE",
     "consultant": "BDA",
     "job_no": "549500",                              # optional
     "items": [
       {"sheet": "A2.201", "title": "BASEMENT PLAN",
        "rev": "C",                                   # exactly as the title block shows
        "grid": "C",                                  # optional short label for the grid (<=4 chars)
        "status": "FOR CONSTRUCTION"}                 # issue status stamped on the document
     ]}
  ]
}

Exit codes: 0 ok, 1 validation failure or conflicts that stopped the write, 2 usage error.
"""
from __future__ import annotations

import copy
import datetime as dt
import json
import re
import sys

try:
    import openpyxl
    from openpyxl.styles import Border, Font, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install --break-system-packages openpyxl")

SHEET = "TRANSMITTAL"
NAVY = "FF0F1721"
GOLD = "FF9D7B5B"
HDR_DD, HDR_MM, HDR_YY = 9, 10, 11
DATA_START = 12
SEP = "   ·   "

HEADER_CELLS = {
    "client": "B3", "project": "B4", "address": "B5", "issued_to": "B6", "status": "B7",
    "tx_no": "K3", "date": "K4", "project_no": "K5", "reason": "K6", "prepared_by": "K7",
}

# Band order follows the 06_Drawings & Specifications sub-folders of the BDM
# New Job Folder. Anything not listed is placed after, in the order first seen.
DISCIPLINE_ORDER = [
    "ARCHITECTURE", "INTERIOR DESIGN", "STRUCTURAL ENGINEERING", "STRUCTURAL",
    "WIND ENGINEERING", "POOL ENGINEERING", "CIVIL ENGINEERING", "CIVIL", "ELECTRICAL",
    "MECHANICAL", "HYDRAULIC", "FIRE SPRINKLERS", "LANDSCAPE", "TRAFFIC", "ACOUSTIC",
    "ENERGY EFFICIENCY", "BUILDING CERTIFIER", "SURVEY", "SURVEYOR", "GEOTECHNICAL",
    "FIRE ENGINEER",
]


# --------------------------------------------------------------------------- helpers
def norm(s) -> str:
    return re.sub(r"\s+", "", str(s or "")).upper()


def natural_key(s: str):
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", norm(s))]


def is_placeholder(v) -> bool:
    return isinstance(v, str) and "{{" in v


def sval(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def layout(ws):
    """Find the date columns, CURRENT ISSUE / STATUS columns and the footer start."""
    cur = stat = None
    for c in range(1, ws.max_column + 1):
        v = sval(ws.cell(HDR_DD, c).value)
        if v and v.upper() == "CURRENT ISSUE":
            cur, stat = c, c + 1
            break
    if not cur:
        raise SystemExit("Not a Form 207 sheet: no CURRENT ISSUE heading in row 9")
    date_cols = list(range(3, cur))
    footer = None
    for r in range(DATA_START, ws.max_row + 1):
        if norm(ws.cell(r, 1).value) == "DISTRIBUTION":
            footer = r
            break
    if footer is None:
        raise SystemExit("Not a Form 207 sheet: no DISTRIBUTION block below the register")
    return date_cols, cur, stat, footer


def read_dates(ws, date_cols):
    dates = []
    for c in date_cols:
        d, m, y = (sval(ws.cell(r, c).value) for r in (HDR_DD, HDR_MM, HDR_YY))
        if d and m and y and d.isdigit() and m.isdigit() and y.isdigit():
            yy = int(y) + (2000 if int(y) < 100 else 0)
            dates.append((c, dt.date(yy, int(m), int(d)).isoformat()))
        else:
            dates.append((c, None))
    return dates


def band_merged(ws, r, last_col) -> bool:
    for m in ws.merged_cells.ranges:
        if m.min_row == r == m.max_row and m.min_col == 1 and m.max_col == last_col:
            return True
    return False


def parse_band(label: str):
    parts = [p.strip() for p in re.split(r"\s+·\s+", label)]
    disc = parts[0] if parts else label
    cons = parts[1] if len(parts) > 1 else ""
    job = ""
    if len(parts) > 2:
        job = re.sub(r"(?i)^job\s*no\.?\s*", "", parts[2]).strip()
    return disc, cons, job


def band_label(b) -> str:
    s = b["discipline"].upper()
    if b.get("consultant"):
        s += SEP + b["consultant"]
    if b.get("job_no"):
        s += SEP + "JOB No. " + b["job_no"]
    return s


# --------------------------------------------------------------------------- read
def read(path):
    wb = openpyxl.load_workbook(path)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"{path}: no {SHEET} sheet — not a Form 207 register")
    ws = wb[SHEET]
    date_cols, cur, stat, footer = layout(ws)
    dates = read_dates(ws, date_cols)
    model = {
        "header": {k: sval(ws[a].value) for k, a in HEADER_CELLS.items()},
        "dates": [d for _, d in dates if d],
        "capacity": len(date_cols),
        "bands": [],
    }
    col_date = {c: d for c, d in dates if d}
    band = None
    for r in range(DATA_START, footer):
        a, b = sval(ws.cell(r, 1).value), sval(ws.cell(r, 2).value)
        if a is None and b is None:
            continue
        if band_merged(ws, r, stat) or (a and b is None and "·" in a):
            disc, cons, job = parse_band(a or "")
            band = {"discipline": disc, "consultant": cons, "job_no": job, "rows": [],
                    "placeholder": is_placeholder(a)}
            model["bands"].append(band)
            continue
        if band is None:
            band = {"discipline": "UNASSIGNED", "consultant": "", "job_no": "", "rows": []}
            model["bands"].append(band)
        grid = {}
        for c in date_cols:
            v = sval(ws.cell(r, c).value)
            if v is not None and c in col_date:
                grid[col_date[c]] = v
        band["rows"].append({
            "sheet": a or "", "title": b or "", "grid": grid,
            "current": sval(ws.cell(r, cur).value) or "",
            "status": sval(ws.cell(r, stat).value) or "",
        })
    return wb, ws, model


# --------------------------------------------------------------------------- apply
def find_band(model, disc, cons):
    for b in model["bands"]:
        if norm(b["discipline"]) == norm(disc) and norm(b["consultant"]) == norm(cons):
            return b
    return None


def insert_band(model, band):
    order = {d: i for i, d in enumerate(DISCIPLINE_ORDER)}
    rank = order.get(band["discipline"].upper(), 999)
    pos = len(model["bands"])
    for i, b in enumerate(model["bands"]):
        if order.get(b["discipline"].upper(), 999) > rank:
            pos = i
            break
    model["bands"].insert(pos, band)


def recompute_current(row, full_labels):
    """CURRENT ISSUE is the revision on the latest-dated issue -- by date, never by
    comparing labels (P3 -> A, 13W -> 17 and 01 -> C all occur in practice)."""
    if not row["grid"]:
        return
    latest = max(row["grid"])
    full = full_labels.get((norm(row["sheet"]), latest))
    if full:
        row["current"] = full
    elif not (row["current"] and norm(row["grid"][latest]) in norm(row["current"])):
        row["current"] = row["grid"][latest]


def apply(model, changes):
    report = {"added": [], "updated": [], "already": [], "older_issue": [],
              "conflicts": [], "title_changed": [], "dropped_dates": [], "new_bands": []}
    if changes.get("placeholders"):
        model["placeholders"] = changes["placeholders"]
    for k, v in (changes.get("header") or {}).items():
        if k not in HEADER_CELLS:
            raise SystemExit(f"Unknown header key {k!r}; use one of {sorted(HEADER_CELLS)}")
        model["header"][k] = v
    # drop template sample content the first time a register is used
    for b in list(model["bands"]):
        if b.get("placeholder"):
            model["bands"].remove(b)

    full_labels = {}
    for issue in changes.get("received", []):
        date = dt.date.fromisoformat(issue["date"]).isoformat()
        band = find_band(model, issue["discipline"], issue.get("consultant", ""))
        if band is None:
            band = {"discipline": issue["discipline"].upper(), "consultant": issue.get("consultant", ""),
                    "job_no": issue.get("job_no", ""), "rows": []}
            insert_band(model, band)
            report["new_bands"].append(band_label(band))
        elif issue.get("job_no") and not band.get("job_no"):
            band["job_no"] = issue["job_no"]
        if date not in model["dates"]:
            model["dates"].append(date)
            model["dates"].sort()
        for it in issue["items"]:
            rev = str(it["rev"]).strip()
            grid_label = str(it.get("grid") or rev).strip()
            if len(grid_label) > 4:
                d = re.sub(r"\D", "", grid_label)
                grid_label = d[-3:] if d else grid_label[:4]
            key = norm(it["sheet"])
            row = next((r for r in band["rows"] if norm(r["sheet"]) == key), None)
            tag = f'{band["discipline"]} {it["sheet"]} rev {rev} ({date})'
            if row is None:
                row = {"sheet": it["sheet"].strip(), "title": (it.get("title") or "").strip().upper(),
                       "grid": {}, "current": "", "status": ""}
                band["rows"].append(row)
                band["rows"].sort(key=lambda r: natural_key(r["sheet"]))
                report["added"].append(tag)
            elif date in row["grid"]:
                if norm(row["grid"][date]) == norm(grid_label):
                    report["already"].append(tag)
                    continue
                report["conflicts"].append(f'{tag}: register already shows {row["grid"][date]} on that date')
                continue
            else:
                if row["grid"] and date < max(row["grid"]):
                    report["older_issue"].append(f'{tag}: older than the current issue {row["current"]}; grid only')
                else:
                    report["updated"].append(f'{tag} (was {row["current"] or "-"})')
                t = (it.get("title") or "").strip().upper()
                if t and row["title"] and norm(t) != norm(row["title"]):
                    report["title_changed"].append(f'{it["sheet"]}: "{row["title"]}" -> "{t}"')
                    if not (row["grid"] and date < max(row["grid"])):
                        row["title"] = t
            row["grid"][date] = grid_label
            full_labels[(key, date)] = rev
            if date >= max(row["grid"]):
                row["status"] = (it.get("status") or row["status"] or "").strip().upper()
            recompute_current(row, full_labels)

    # roll: keep the most recent `capacity` issue dates; history stays in the _SS copies
    cap = model["capacity"]
    if len(model["dates"]) > cap:
        drop = model["dates"][: len(model["dates"]) - cap]
        model["dates"] = model["dates"][len(drop):]
        report["dropped_dates"] = drop
        for b in model["bands"]:
            for r in b["rows"]:
                for d in drop:
                    r["grid"].pop(d, None)
    return report


# --------------------------------------------------------------------------- render
def navy(cell_style_src, target):
    target._style = copy.copy(cell_style_src._style)
    f = copy.copy(target.font)
    target.font = Font(name=f.name, sz=f.sz, b=f.b, i=f.i, color=NAVY)


def render(wb, ws, model):
    date_cols, cur, stat, footer = layout(ws)
    last_col = stat

    # prototypes taken from the sheet itself so the form's own styling is used
    band_proto = band_row_proto = None
    for r in range(DATA_START, footer):
        if band_proto is None and band_merged(ws, r, last_col):
            band_proto = r
        elif band_proto is not None and sval(ws.cell(r, 1).value):
            band_row_proto = r
            break
    if band_proto is None or band_row_proto is None:
        raise SystemExit("Cannot find a band row and a document row to take styles from")
    proto_band = copy.copy(ws.cell(band_proto, 1)._style)
    proto_cells = {c: copy.copy(ws.cell(band_row_proto, c)) for c in range(1, last_col + 1)}
    proto_styles = {c: copy.copy(ws.cell(band_row_proto, c)._style) for c in range(1, last_col + 1)}
    thin = Side(style="thin", color=None)

    # snapshot the footer block (DISTRIBUTION onward)
    foot = []
    for r in range(footer, ws.max_row + 1):
        foot.append({
            "h": ws.row_dimensions[r].height,
            "cells": [(c, ws.cell(r, c).value, copy.copy(ws.cell(r, c)._style)) for c in range(1, last_col + 1)],
        })
    foot_merges = [(m.min_row - footer, m.min_col, m.max_row - footer, m.max_col)
                   for m in ws.merged_cells.ranges if m.min_row >= footer]
    for m in [m for m in ws.merged_cells.ranges if m.min_row >= DATA_START]:
        ws.unmerge_cells(str(m))

    # clear everything below the date header
    blank = copy.copy(ws.cell(8, 1)._style)
    for r in range(DATA_START, ws.max_row + 1):
        ws.row_dimensions[r].height = None
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            cell.value = None
            cell._style = copy.copy(blank)

    # header fields
    for k, a in HEADER_CELLS.items():
        v = model["header"].get(k)
        if v is not None:
            ws[a].value = v
            if not is_placeholder(v):
                f = ws[a].font
                ws[a].font = Font(name=f.name, sz=f.sz, b=f.b, i=f.i, color=NAVY)

    # date header
    col_of = {}
    for i, c in enumerate(date_cols):
        d = model["dates"][i] if i < len(model["dates"]) else None
        vals = (None, None, None)
        if d:
            y, m, dd = d.split("-")
            vals = (dd, m, y[2:])
            col_of[d] = c
        for r, v in zip((HDR_DD, HDR_MM, HDR_YY), vals):
            cell = ws.cell(r, c)
            cell.value = v
            f = cell.font
            cell.font = Font(name=f.name, sz=f.sz, b=f.b, i=f.i, color=NAVY)

    # bands and rows
    r = DATA_START
    for b in model["bands"]:
        cell = ws.cell(r, 1)
        cell.value = band_label(b)
        cell._style = copy.copy(proto_band)
        f = cell.font
        cell.font = Font(name=f.name, sz=f.sz, b=f.b, i=f.i, color=NAVY)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)
        r += 1
        for i, row in enumerate(b["rows"]):
            last = i == len(b["rows"]) - 1
            values = {1: row["sheet"], 2: row["title"], cur: row["current"] or None, stat: row["status"] or None}
            for d, v in row["grid"].items():
                if d in col_of:
                    values[col_of[d]] = v
            for c in range(1, last_col + 1):
                cell = ws.cell(r, c)
                cell._style = copy.copy(proto_styles[c])
                f = cell.font
                cell.font = Font(name=f.name, sz=f.sz, b=f.b, i=f.i, color=NAVY)
                if last:
                    bd = cell.border
                    cell.border = Border(left=bd.left, right=bd.right, top=bd.top, bottom=thin)
                cell.value = values.get(c)
            r += 1
    r += 1  # spacer

    # footer back in place
    new_footer = r
    for k, fr in enumerate(foot):
        ws.row_dimensions[new_footer + k].height = fr["h"]
        for c, v, st in fr["cells"]:
            cell = ws.cell(new_footer + k, c)
            cell.value = v
            cell._style = st
    for (r1, c1, r2, c2) in foot_merges:
        ws.merge_cells(start_row=new_footer + r1, start_column=c1, end_row=new_footer + r2, end_column=c2)
    last_row = new_footer + len(foot) - 1
    # trim anything left below the footer from a longer previous version
    if ws.max_row > last_row:
        ws.delete_rows(last_row + 1, ws.max_row - last_row)

    # replace any remaining form placeholders the caller has supplied text for
    repl = model.get("placeholders") or {}
    if repl:
        for row in ws.iter_rows(min_row=1, max_row=last_row, max_col=last_col):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip() in repl:
                    cell.value = repl[cell.value.strip()]
                    f = cell.font
                    cell.font = Font(name=f.name, sz=f.sz, b=f.b, i=f.i, color=NAVY)

    ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"
    ws.print_title_rows = f"{HDR_DD}:{HDR_YY}"
    ws.freeze_panes = f"C{DATA_START}"
    return last_row


# --------------------------------------------------------------------------- validate
def validate(path):
    problems, warnings = [], []
    try:
        wb, ws, model = read(path)
    except SystemExit as e:
        return [str(e)], []
    date_cols, cur, stat, footer = layout(ws)
    if model["dates"] != sorted(set(model["dates"])):
        problems.append("Issue dates in the header are not unique and ascending")
    for b in model["bands"]:
        seen = set()
        for r in b["rows"]:
            k = norm(r["sheet"])
            if k in seen:
                problems.append(f'{b["discipline"]}: sheet {r["sheet"]} appears twice — one row per document')
            seen.add(k)
            if r["grid"]:
                latest = r["grid"][max(r["grid"])]
                if not r["current"]:
                    problems.append(f'{r["sheet"]}: CURRENT ISSUE is blank')
                elif norm(latest) not in norm(r["current"]) and norm(r["current"]) not in norm(latest):
                    problems.append(f'{r["sheet"]}: CURRENT ISSUE {r["current"]} does not match latest grid entry {latest}')
            elif not r["current"]:
                warnings.append(f'{r["sheet"]}: no issue recorded in the grid and no CURRENT ISSUE')
            if is_placeholder(r["sheet"]) or is_placeholder(r["title"]):
                problems.append(f'{r["sheet"]}: placeholder text left in a document row')
    for k, v in model["header"].items():
        if is_placeholder(v):
            warnings.append(f"Header field {k} still carries the placeholder {v}")
    left = sorted({c.value.strip() for row in ws.iter_rows() for c in row
                   if isinstance(c.value, str) and "{{" in c.value})
    if left:
        warnings.append("Placeholders still on the sheet: " + ", ".join(left))
    pa = ws.print_area
    if not pa or not str(pa).endswith(str(ws.max_row)):
        problems.append(f"Print area {pa} does not reach the last row {ws.max_row}")
    for r in range(DATA_START, footer):
        a = sval(ws.cell(r, 1).value)
        if a and sval(ws.cell(r, 2).value) is None and "·" in a and not band_merged(ws, r, stat):
            problems.append(f"Band row {r} is not merged across the sheet")
    return problems, warnings


# --------------------------------------------------------------------------- cli
def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 2
    cmd = argv[1]
    if cmd == "read" and len(argv) == 3:
        _, _, model = read(argv[2])
        print(json.dumps(model, indent=1))
        return 0
    if cmd in ("apply", "new") and len(argv) == 5:
        src, spec, out = argv[2], argv[3], argv[4]
        if out == src:
            print("Refusing to overwrite the input — save to a new filename")
            return 2
        with open(spec, encoding="utf-8") as fh:
            changes = json.load(fh)
        wb, ws, model = read(src)
        if cmd == "new":
            if not changes.get("received"):
                print("A new register is created at the first receipt — 'received' must not be empty")
                return 2
            if model["dates"] or any(not b.get("placeholder") for b in model["bands"]):
                print("'new' takes the blank Form 207 template, not a live register — use 'apply'")
                return 2
        report = apply(model, changes)
        if report["conflicts"]:
            print(json.dumps(report, indent=1))
            print("STOPPED — conflicts must be resolved by a person; nothing written")
            return 1
        render(wb, ws, model)
        wb.save(out)
        problems, warnings = validate(out)
        report["validate"] = {"problems": problems, "warnings": warnings}
        print(json.dumps(report, indent=1))
        return 1 if problems else 0
    if cmd == "validate" and len(argv) == 3:
        problems, warnings = validate(argv[2])
        print(json.dumps({"problems": problems, "warnings": warnings}, indent=1))
        return 1 if problems else 0
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv))
